import argparse
import hashlib
import json
import re
from datetime import date, datetime
from pathlib import Path

import chromadb
from sentence_transformers import SentenceTransformer

from servers.schemes_kb.extract_documents import build_records, extract_document
from servers.schemes_kb.models import Scheme
from servers.schemes_kb.scheme_loading import load_schemes
from servers.schemes_kb.preview_chunks import PROJECT_ROOT, make_chunks

DB_PATH = PROJECT_ROOT / "data" / "chroma"
COLLECTION_NAME = "schemes_bge_m3_body_v1"
REVIEW_COLLECTION_NAME = "schemes_bge_m3_body_v1_review"
MODEL_NAME = "BAAI/bge-m3"




def schemes_from_excel(path):
    from openpyxl import load_workbook

    # Read the review workbook without changing it.
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        def rows(sheet):
            values = workbook[sheet].iter_rows(min_row=5, values_only=True)
            # Row 5 contains headers; map each following row to those column names.
            headers = next(values)
            return [dict(zip(headers, row)) for row in values if any(v is not None for v in row)]

        overview = rows("Scheme review")
        passages = rows("Source passages")
        eligibility = rows("Eligibility")
        applications = rows("Application steps")
        documents = rows("Applicant documents")
        links = rows("Documents and links")
        required = [(overview, "ID"), (passages, "Scheme ID"),
                    (applications, "Scheme ID"), (documents, "Scheme ID"), (links, "Scheme IDs")]
        for records, key in required:
            if not records or key not in records[0]:
                raise ValueError(f"Missing Excel rows or header: {key}")

        # Rebuild a section from approved rows, including text split across multiple rows.
        def section(records, text_key, path_key, status_key, part_key, heading_key, channel_key=None):
            groups = {}
            for row in records:
                if row.get(status_key) != "Approved":
                    raise ValueError("Approved scheme has an unapproved related row")
                path = row.get(path_key)
                if not path:
                    if row.get(text_key) in ("Not stated in captured section", "Not stated in the captured source. Do not infer no restrictions."):
                        continue
                    raise ValueError("Missing source path in Excel")
                text = row.get(text_key)
                if not isinstance(text, str) or not text.strip():
                    raise ValueError(f"Missing source text: {path}")
                part = str(row.get(part_key) or "1")
                if "/" in part:
                    number, total = map(int, part.split("/"))
                else:
                    number, total = int(part), int(row.get("Parts") or 1)
                # Keep different source passages and application channels separate.
                key = (path, row.get(channel_key) if channel_key else None)
                group = groups.setdefault(key, {"parts": {}, "total": total, "heading": row.get(heading_key)})
                if total != group["total"] or number in group["parts"]:
                    raise ValueError(f"Duplicate or inconsistent continuation: {path}")
                group["parts"][number] = text
            items = []
            for (path, channel), group in groups.items():
                # Require every part before joining; otherwise we could lose scheme conditions.
                if set(group["parts"]) != set(range(1, group["total"] + 1)):
                    raise ValueError(f"Missing continuation part: {path}")
                items.append({"text": "".join(group["parts"][i] for i in range(1, group["total"] + 1)),
                              "source_path": path, "source_url": source_url,
                              "heading": group["heading"] or None, "channel": channel})
            return {"capture_status": "captured" if items else "not_in_source", "items": items}

        schemes = []
        seen = set()
        for row in overview:
            sid = row["ID"]
            if sid in seen:
                raise ValueError(f"Duplicate Excel scheme ID: {sid}")
            seen.add(sid)
            # A scheme and its related content must be approved before export.
            if row.get("Review status") != "Approved":
                print(f"Skipped Excel scheme {sid}: {row.get('Review status')}")
                continue
            slug = row.get("myScheme slug")
            if not isinstance(slug, str) or not re.fullmatch(r"[a-z0-9][a-z0-9-]*", slug):
                raise ValueError(f"Invalid or missing myScheme slug: {sid}")
            source_url = f"https://www.myscheme.gov.in/schemes/{slug}"
            for related in eligibility:
                if related.get("Scheme ID") == sid and related.get("Review status") != "Approved":
                    raise ValueError(f"Unapproved eligibility row: {sid}")
            source_rows = [r for r in passages if r.get("Scheme ID") == sid]
            if not source_rows:
                raise ValueError(f"No source passages for {sid}")
            if any(r.get("Review state") != "Approved" for r in source_rows):
                raise ValueError(f"Unapproved source passage: {sid}")
            # Use Source passages for the main scheme sections.
            sections = {}
            for name, field in [("Details", "details"), ("Benefits", "benefits"),
                                ("Eligibility", "eligibility"), ("Exclusions", "exclusions")]:
                sections[field] = section([r for r in source_rows if r.get("Section") == name],
                                          "Source text", "Original source path", "Review state",
                                          "Part", "Source heading")
            sections["application_process"] = section(
                [r for r in applications if r.get("Scheme ID") == sid], "Step / process text",
                "Original source path", "Review status", "Continuation part", "Source phase / heading", "Source mode")
            # Document-review rows supply text; passage IDs restore source paths and part numbers.
            document_rows = []
            part_counts = {}
            for related in documents:
                if related.get("Scheme ID") != sid:
                    continue
                if related.get("Review status") != "Approved":
                    raise ValueError(f"Unapproved applicant-document row: {sid}")
                if related.get("Applicant document / source text") == "Not stated in captured section":
                    continue
                pid = related.get("Passage / row ID")
                part_counts[pid] = part_counts.get(pid, 0) + 1
                matches = [r for r in source_rows if r.get("Passage ID") == pid
                           and int(r.get("Part") or 1) == part_counts[pid]]
                if len(matches) != 1:
                    raise ValueError(f"Missing document source pointer: {sid} / {pid}")
                record = dict(matches[0])
                record["Source text"] = related["Applicant document / source text"]
                record["Source heading"] = related.get("Source context")
                document_rows.append(record)
            sections["required_documents"] = section(document_rows, "Source text", "Original source path",
                                                     "Review state", "Part", "Source heading")
            # Detect review rows accidentally removed from the application or document tabs.
            for name, field in [("Application process", "application_process"), ("Applicant documents", "required_documents")]:
                expected = {r["Original source path"] for r in source_rows if r.get("Section") == name}
                actual = {item["source_path"] for item in sections[field]["items"]}
                if actual != expected:
                    raise ValueError(f"Review-tab rows do not match source paths: {sid} / {name}")
            collected = row.get("Research date")
            if isinstance(collected, datetime):
                collected = collected.date()
            if not isinstance(collected, date):
                raise ValueError(f"Missing or invalid research date: {sid}")
            # Capture published link titles and URLs; this does not verify their destinations.
            resources = []
            for link in links:
                if link.get("Scheme IDs") != sid or link.get("Document type") == "Official scheme page":
                    continue
                if "Review status: Approved" not in (link.get("Reviewer notes") or ""):
                    raise ValueError(f"Unapproved link for {sid}")
                url = link.get("URL")
                title = link.get("Document / page title") or "Source link"
                kind = "other"
                if "application" in title.lower() and "form" in title.lower():
                    kind = "application_form"
                elif "guideline" in title.lower():
                    kind = "guidelines"
                resources.append({"title": title, "url": url, "source_url": source_url,
                                  "source_path": link.get("Where to look / source path"), "type": kind,
                                  "format": "pdf" if re.search(r"\.pdf(?:[?#]|$)", url or "", re.I) else "webpage",
                                  "verification_status": "pending"})
            availability = {"Available": "open", "Open": "open", "Closed": "closed",
                            "Not verified": "unverified", "Unverified": "unverified"}
            value = row.get("Application availability")
            if value not in availability:
                raise ValueError(f"Unknown availability for {sid}: {value}")
            # Pydantic checks that the record matches our shared scheme schema.
            schemes.append(Scheme.model_validate({
                "scheme_id": sid, "slug": slug, "name": row["Current source name / alias"] or row["Scheme"],
                "government_level": row["Government"], "geography": row["Geography"],
                "source_url": source_url, "collected_on": collected, "review_status": "approved",
                "availability_status": availability[value], "sections": sections, "links": resources,
            }))
        if not schemes:
            raise ValueError("Excel contains no approved schemes; no files or database were changed.")
        return schemes
    finally:
        workbook.close()


def export_excel_schemes(path, directory):
    # Validate the entire workbook before writing any scheme files.
    schemes = schemes_from_excel(path)
    existing = {s.scheme_id: s for s in load_schemes(directory)} if directory.exists() else {}
    exports = []
    for scheme in schemes:
        target = directory / f"{scheme.slug}.json"
        if target.exists():
            previous = Scheme.model_validate_json(target.read_text(encoding="utf-8"))
            if previous.scheme_id != scheme.scheme_id:
                raise ValueError(f"Filename belongs to another scheme: {target}")
        previous = existing.get(scheme.scheme_id)
        if previous and previous.slug != scheme.slug:
            raise ValueError(f"Scheme slug changed: {scheme.scheme_id}; reconcile existing JSON first.")
        # Keep existing PDF registrations when the exported link URL still matches.
        for link in scheme.links:
            if previous:
                old = next((item for item in previous.links if item.url == link.url), None)
                if old:
                    link.local_file = old.local_file
                    link.verification_status = old.verification_status
        exports.append((target, scheme.model_dump_json(indent=2) + "\n"))
    if len({target for target, _ in exports}) != len(exports):
        raise ValueError("Duplicate Excel slugs; no JSON files were written.")
    directory.mkdir(parents=True, exist_ok=True)
    for target, content in exports:
        # Replace each file only after its complete JSON has been written.
        temporary = target.with_suffix(".json.tmp")
        temporary.write_text(content, encoding="utf-8")
        temporary.replace(target)
    print(f"Exported {len(exports)} validated schemes to {directory}")
    return schemes


# Chroma stores nested link lists as JSON strings in metadata.
def link_metadata(scheme):
    links = [link.model_dump(mode="json", exclude={"local_file"})
             for link in getattr(scheme, "links", [])]
    # Source paths identify links from myScheme’s Sources and References section.
    references = [link for link in links if ".references[" in (link.get("source_path") or "")]
    return {"links_json": json.dumps(links, ensure_ascii=False),
            "references_json": json.dumps(references, ensure_ascii=False)}


# PDFs are optional; only registered files with approved extractions are selected.
def prepare_pdfs(schemes, model):
    selected = {}
    for scheme in schemes:
        if scheme.review_status != "approved":
            continue
        folder = (PROJECT_ROOT / "data/pdfs" / scheme.scheme_id).resolve()
        registered = set()
        for link in scheme.links:
            if not link.local_file:
                continue
            path = (PROJECT_ROOT / link.local_file).resolve()
            if not path.is_relative_to(folder) or path.suffix.lower() != ".pdf":
                raise ValueError(f"{scheme.scheme_id}: local_file must be a PDF inside {folder}")
            registered.add(path)
            extracted = extract_document(scheme, path, str(link.url), link.title,
                                         collected_on=scheme.collected_on)
            document = json.loads(extracted.read_text(encoding="utf-8"))
            if document["review_status"] != "approved":
                print(f"Pending PDF review: {extracted}")
                continue
            # Confirm the reviewed extraction belongs to this exact PDF and scheme.
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            raw = (PROJECT_ROOT / document["file_path"]).resolve()
            if (document["scheme_id"] != scheme.scheme_id or
                    document["file_sha256"] != digest or
                    document["document_id"] != f"{scheme.scheme_id}-{digest}" or
                    document["source_url"] != str(link.url) or
                    document.get("extraction_version") != "tables_v1" or
                    not raw.is_relative_to((PROJECT_ROOT / "data/raw").resolve()) or
                    hashlib.sha256(raw.read_bytes()).hexdigest() != digest):
                raise ValueError(f"PDF provenance differs: {extracted}; re-extract and review.")
            if document.get("ocr_review_pages") or any(p["needs_ocr_review"] for p in document["pages"]):
                raise ValueError(f"Resolve OCR review flags before importing {extracted}")
            records = build_records(document, model.tokenizer)
            if not records or any(len(model.tokenizer.encode(r["text"])) > model.max_seq_length for r in records):
                raise ValueError(f"Empty PDF or oversized table: {extracted}")
            selected[document["document_id"]] = (scheme, document, records)
        if folder.is_dir():
            for path in folder.rglob("*"):
                if path.suffix.lower() == ".pdf" and path.resolve() not in registered:
                    raise ValueError(f"Register this PDF in scheme links with local_file and official URL: {path}")
    return selected


def import_pdfs(client, selected, model):
    collection = client.get_or_create_collection(
        "scheme_pdf_review_v1", embedding_function=None,
        configuration={"hnsw": {"space": "cosine"}},
        metadata={"model": MODEL_NAME, "owner": "import_schemes", "input_strategy": "pdf_tables_v1"},
    )
    if (collection.metadata or {}).get("model") != MODEL_NAME or (collection.metadata or {}).get("input_strategy") != "pdf_tables_v1":
        raise ValueError("PDF collection uses another embedding strategy.")
    # Track this run’s chunk IDs so outdated approved PDF chunks can be removed.
    current_ids = set()
    # Embed and upsert in batches: existing IDs are updated, new IDs are inserted.
    for scheme, document, records in selected.values():
        for start in range(0, len(records), 16):
            batch = records[start:start + 16]
            texts = [r["text"] for r in batch]
            vectors = model.encode(texts, batch_size=2, normalize_embeddings=True)
            collection.upsert(ids=[r["id"] for r in batch], documents=texts,
                              embeddings=vectors.tolist(), metadatas=[{
                "scheme_id": scheme.scheme_id, "document_id": document["document_id"],
                "page": r["page"], "kind": r["kind"], "source_url": document["source_url"],
                "review_status": "approved", "availability_status": scheme.availability_status,
                "source_type": "pdf", "file_sha256": document["file_sha256"],
                "owner": "import_schemes",
                **link_metadata(scheme),
            } for r in batch])
            current_ids.update(r["id"] for r in batch)
    existing = collection.get(include=["metadatas"])
    stale = [key for key, meta in zip(existing["ids"], existing["metadatas"])
             if key not in current_ids and
             (meta.get("owner") == "import_schemes" or meta.get("review_status") == "approved")]
    for start in range(0, len(stale), 500):
        collection.delete(ids=stale[start:start + 500])
    print(f"Saved {len(current_ids)} PDF chunks from {len(selected)} approved PDFs; removed {len(stale)} stale chunks.")


def main():
    parser = argparse.ArgumentParser(description="Bulk import approved scheme pages and reviewed local PDFs into ChromaDB.")
    parser.add_argument("--review", action="store_true", help="Include drafts in a separate review collection.")
    parser.add_argument("--from-excel", nargs="?", const=str(PROJECT_ROOT / "docs/Scheme-Saathi-60-Schemes-Schema-Review.xlsx"),
                        type=Path, help="Convert approved Excel rows to scheme JSON before import; optionally provide a workbook path.")
    parser.add_argument("--export-only", action="store_true", help="Write validated JSON without loading embeddings or changing ChromaDB.")
    args = parser.parse_args()
    if args.export_only and not args.from_excel:
        parser.error("--export-only requires --from-excel")
    # Export-only stops before loading the embedding model or opening ChromaDB.
    if args.from_excel:
        try:
            export_excel_schemes(args.from_excel, PROJECT_ROOT / "data/schemes")
        except (ValueError, KeyError, OSError) as error:
            parser.error(str(error))
        if args.export_only:
            return
    # Draft scheme pages use a separate collection when --review is supplied.
    collection_name = REVIEW_COLLECTION_NAME if args.review else COLLECTION_NAME
    schemes = []
    seen = set()
    for scheme in load_schemes(PROJECT_ROOT / "data" / "schemes"):
        if scheme.sample_only:
            continue
        if scheme.scheme_id in seen:
            raise ValueError(f"Duplicate scheme ID: {scheme.scheme_id}")
        seen.add(scheme.scheme_id)
        if scheme.review_status == "approved" or (args.review and scheme.review_status == "draft"):
            schemes.append(scheme)
        else:
            print(f"Skipped {scheme.scheme_id}: {scheme.review_status}")

    pdfs = {}
    chunks = []
    metadata = []
    if schemes:
        # Reuse the downloaded model locally; normalized vectors support cosine search.
        model = SentenceTransformer(MODEL_NAME, device="cpu", local_files_only=True)
        if not args.review:
            pdfs = prepare_pdfs(schemes, model)
        # Split each scheme into overlapping chunks and retain citation metadata.
        for scheme in schemes:
            scheme_hash = hashlib.sha256(scheme.model_dump_json().encode()).hexdigest()
            for chunk in make_chunks(scheme, model.tokenizer):
                chunks.append(chunk)
                metadata.append({
                    "scheme_id": scheme.scheme_id,
                    "scheme_name": scheme.name,
                    **link_metadata(scheme),
                    "section": chunk["section"],
                    "government_level": scheme.government_level,
                    "geography": scheme.geography,
                    "collected_on": scheme.collected_on.isoformat(),
                    "review_status": scheme.review_status,
                    "availability_status": scheme.availability_status,
                    "capture_status": chunk["capture_status"],
                    "source_url": str(scheme.source_url),
                    "citations_json": json.dumps(chunk["citations"], ensure_ascii=False),
                    "context_json": json.dumps(chunk["context"], ensure_ascii=False),
                    "scheme_hash": scheme_hash,
                    "model": MODEL_NAME,
                    "input_strategy": "body_v1",
                    "token_start": chunk["token_start"],
                    "token_end": chunk["token_end"],
                })

    # PersistentClient saves the collections on disk for later retrieval.
    client = chromadb.PersistentClient(path=str(DB_PATH))
    collection = client.get_or_create_collection(
        name=collection_name,
        embedding_function=None,
        configuration={"hnsw": {"space": "cosine"}},
        metadata={"model": MODEL_NAME, "input_strategy": "body_v1", "owner": "import_schemes"},
    )
    if collection.metadata != {"model": MODEL_NAME, "input_strategy": "body_v1", "owner": "import_schemes"}:
        raise ValueError("Collection configuration differs; use a separate collection.")

    # Write chunks, embeddings and metadata together using stable chunk IDs.
    for start in range(0, len(chunks), 32):
        batch = chunks[start:start + 32]
        texts = [chunk["text"] for chunk in batch]
        vectors = model.encode(texts, batch_size=2, normalize_embeddings=True)
        collection.upsert(
            ids=[chunk["preview_id"] for chunk in batch],
            documents=texts,
            embeddings=vectors.tolist(),
            metadatas=metadata[start:start + 32],
        )

    # Each collection mirrors its selected corpus; obsolete chunks are removed after upserts.
    current_ids = {chunk["preview_id"] for chunk in chunks}
    existing_ids = collection.get(include=[])["ids"]
    stale_ids = [key for key in existing_ids if key not in current_ids]
    for start in range(0, len(stale_ids), 500):
        collection.delete(ids=stale_ids[start:start + 500])

    if not args.review:
        import_pdfs(client, pdfs, model if schemes else None)

    print(f"Saved {len(chunks)} chunks from {len(schemes)} selected schemes.")
    print(f"Removed {len(stale_ids)} stale chunks. Collection count: {collection.count()}")
    print(f"Database: {DB_PATH}")
    print(f"Collection: {collection_name}")


if __name__ == "__main__":
    main()
