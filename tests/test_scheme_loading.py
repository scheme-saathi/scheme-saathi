import json
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch
import hashlib
from pathlib import Path

from pydantic import ValidationError

from servers.schemes_kb.scheme_loading import load_schemes

PROJECT_ROOT = Path(__file__).resolve().parents[1]


class SchemeLoadingChecks(unittest.TestCase):
    def test_real_scheme_and_synthetic_fixture_can_coexist(self):
        with tempfile.TemporaryDirectory() as folder:
            directory = Path(folder)
            for name in ['rkvp.json', 'SYNTH-PENSION-001.json']:
                (directory / name).write_bytes((PROJECT_ROOT / 'data/schemes' / name).read_bytes())
            records = list(load_schemes(directory))
            self.assertEqual([r.scheme_id for r in records], ['C07'])

    def test_invalid_real_record_is_not_silently_skipped(self):
        with tempfile.TemporaryDirectory() as folder:
            directory = Path(folder)
            (directory / 'broken.json').write_text(json.dumps({'scheme_id': 'C99'}))
            with self.assertRaises(ValidationError):
                list(load_schemes(directory))

    def test_name_alone_does_not_bypass_validation(self):
        with tempfile.TemporaryDirectory() as folder:
            directory = Path(folder)
            (directory / 'SYNTH-broken.json').write_text(json.dumps({'scheme_id': 'SYNTH-broken'}))
            with self.assertRaises(ValidationError):
                list(load_schemes(directory))


class BulkPdfChecks(unittest.TestCase):
    def test_pdf_approval_hash_and_scheme_gate(self):
        from servers.schemes_kb import import_schemes as importer
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            pdf = root / "data/pdfs/C07/guide.pdf"
            pdf.parent.mkdir(parents=True)
            pdf.write_bytes(b"test PDF bytes")
            digest = hashlib.sha256(pdf.read_bytes()).hexdigest()
            raw = root / "data/raw/C07/source.pdf"
            raw.parent.mkdir(parents=True)
            raw.write_bytes(pdf.read_bytes())
            extracted = root / "extracted.json"
            doc = {"scheme_id": "C07", "document_id": f"C07-{digest}",
                   "file_sha256": digest, "file_path": "data/raw/C07/source.pdf",
                   "source_url": "https://example.org/guide.pdf", "extraction_version": "tables_v1",
                   "review_status": "draft", "pages": [{"needs_ocr_review": False}]}
            link = SimpleNamespace(local_file="data/pdfs/C07/guide.pdf", url=doc["source_url"], title="Guide")
            scheme = SimpleNamespace(scheme_id="C07", review_status="approved", links=[link], collected_on=None)
            model = SimpleNamespace(tokenizer=SimpleNamespace(encode=lambda text: [1]), max_seq_length=100)
            with patch.object(importer, "PROJECT_ROOT", root), patch.object(importer, "extract_document", return_value=extracted), patch.object(importer, "build_records", return_value=[{"text": "Evidence"}]):
                extracted.write_text(json.dumps(doc))
                self.assertEqual(importer.prepare_pdfs([scheme], model), {})
                doc["review_status"] = "approved"
                extracted.write_text(json.dumps(doc))
                self.assertEqual(len(importer.prepare_pdfs([scheme], model)), 1)
                raw.write_bytes(b"changed")
                with self.assertRaisesRegex(ValueError, "provenance"):
                    importer.prepare_pdfs([scheme], model)
                scheme.review_status = "draft"
                self.assertEqual(importer.prepare_pdfs([scheme], model), {})

    def test_bulk_upsert_is_repeatable_and_revocation_removes_records(self):
        import chromadb
        import numpy as np
        from servers.schemes_kb.import_schemes import import_pdfs
        with tempfile.TemporaryDirectory() as folder:
            client = chromadb.PersistentClient(path=folder)
            model = SimpleNamespace(encode=lambda texts, **kwargs: np.array([[1., 0.] for _ in texts]))
            scheme = SimpleNamespace(scheme_id="C07", availability_status="open")
            document = {"document_id": "C07-hash", "source_url": "https://example.org/guide.pdf", "file_sha256": "hash"}
            records = [{"id": "C07-hash:page-1", "page": 1, "kind": "prose", "text": "Evidence"}]
            selected = {"C07-hash": (scheme, document, records)}
            import_pdfs(client, selected, model)
            import_pdfs(client, selected, model)
            collection = client.get_collection("scheme_pdf_review_v1")
            self.assertEqual(collection.count(), 1)
            self.assertEqual(collection.get()["metadatas"][0]["review_status"], "approved")
            import_pdfs(client, {}, model)
            self.assertEqual(collection.count(), 0)


if __name__ == '__main__':
    unittest.main()


class ExcelImportChecks(unittest.TestCase):
    def test_export_preserves_text_links_and_pdf_registration(self):
        from servers.schemes_kb.import_schemes import export_excel_schemes, schemes_from_excel
        workbook = PROJECT_ROOT / 'docs/Scheme-Saathi-60-Schemes-Schema-Review.xlsx'
        records = schemes_from_excel(workbook)
        self.assertTrue(records)
        self.assertNotIn('M03', [s.scheme_id for s in records])
        self.assertNotIn('M04', [s.scheme_id for s in records])
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder)
            first = records[0]
            self.assertTrue(first.links)
            first.links[0].local_file = f'data/pdfs/{first.scheme_id}/guide.pdf'
            (destination / f'{first.slug}.json').write_text(first.model_dump_json())
            exported = export_excel_schemes(workbook, destination)
            loaded = list(load_schemes(destination))
            self.assertEqual(len(loaded), len(records))
            self.assertEqual(exported[0].links[0].local_file, first.links[0].local_file)
            for scheme in exported:
                self.assertEqual(scheme.sections, next(s.sections for s in records if s.scheme_id == scheme.scheme_id))
            self.assertEqual(next(s for s in exported if s.scheme_id == 'C11').availability_status, 'closed')

    def test_unapproved_related_row_blocks_all_exports(self):
        from openpyxl import load_workbook
        from servers.schemes_kb.import_schemes import export_excel_schemes
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook = load_workbook(PROJECT_ROOT / 'docs/Scheme-Saathi-60-Schemes-Schema-Review.xlsx')
            sheet = workbook['Eligibility']
            for row in sheet.iter_rows(min_row=6):
                if row[0].value == 'C01':
                    row[7].value = 'Not started'
                    break
            path = root / 'review.xlsx'
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, 'Unapproved eligibility'):
                export_excel_schemes(path, root / 'schemes')
            self.assertFalse((root / 'schemes').exists())

    def test_missing_continuation_blocks_export(self):
        from openpyxl import load_workbook
        from servers.schemes_kb.import_schemes import export_excel_schemes
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            workbook = load_workbook(PROJECT_ROOT / 'docs/Scheme-Saathi-60-Schemes-Schema-Review.xlsx')
            sheet = workbook['Source passages']
            for row in sheet.iter_rows(min_row=6):
                if row[1].value == 'C01' and row[3].value == 'Benefits':
                    row[7].value = 2
                    row[8].value = 2
                    break
            path = root / 'review.xlsx'
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, 'Missing continuation'):
                export_excel_schemes(path, root / 'schemes')
            self.assertFalse((root / 'schemes').exists())
